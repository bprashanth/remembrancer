#!/usr/bin/env python3
"""Positional cell-level scorer — the primary metric for paperroast.

The sibling project steered on a token-MULTISET score, which is blind to
position: a model can emit every value in the wrong cell and still score 0.85.
The product bar is "does the spreadsheet look like the form, with the right
values in the right cells?", so the primary metric here is:

    cell_accuracy = fraction of template cells whose extracted value equals
                    the golden value at that same (row, col) — blanks included.

Both golden and candidate are GRID xlsx files with the same shape as the form:
  sheet "table"  : row 1 = printed column labels, rows 2..N+1 = the N data
                   rows of the template, one column per template column,
                   empty cells left empty.
  sheet "header" : one row per header field: label | value.

Because both sides are written against the SAME template, alignment is by
construction — the scorer never has to guess which row is which. That is the
entire point of owning the form builder.

Reported buckets:
  cell_acc        : all table cells, positional (headline number)
  filled_recall   : golden-nonempty cells read correctly
  blank_acc       : golden-empty cells correctly left empty
  hallucination   : golden-empty cells the candidate invented a value for
  miss            : golden-nonempty cells the candidate left empty
  wrong           : golden-nonempty cells with a different nonempty value
  header_acc      : header fields, matched by label
  per-column      : cell_acc per template column (find weak column types)
  multiset num/word/code F1 : the sibling's comparability metric, secondary

Usage:
    python3 -m core.score golden.xlsx candidate.xlsx [--json]
"""
from __future__ import annotations

import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

import openpyxl

_WS = re.compile(r"\s+")
_NUM = re.compile(r"^[+-]?\d+([.,]\d+)?$")
_TOK = re.compile(r"[^0-9a-z.]+")


# ── normalisation ─────────────────────────────────────────────────
def norm(v) -> str:
    """Canonical form of one cell value for positional comparison."""
    if v is None:
        return ""
    s = str(v).strip()
    if not s:
        return ""
    s = unicodedata.normalize("NFKC", s)
    s = _WS.sub(" ", s)
    # numbers: 23 == 23.0 == "23,0"; keep sign; strip trailing zeros
    t = s.replace(",", ".")
    if _NUM.match(t):
        f = float(t)
        return str(int(f)) if f == int(f) else repr(f)
    return s.casefold()


def cells_equal(g, c) -> bool:
    return norm(g) == norm(c)


# ── grid loading ──────────────────────────────────────────────────
def load_grid(path: str | Path) -> dict:
    """Read a grid xlsx -> {"table": [[...]], "header": [(label, value)]}.

    The table is returned as rows of raw values (None for empty). Trailing
    all-empty rows are NOT trimmed — the template row count is part of the
    contract and the caller compares against the golden's shape.
    """
    wb = openpyxl.load_workbook(str(path), data_only=True)
    out = {"table": [], "header": []}
    if "table" in wb.sheetnames:
        for row in wb["table"].iter_rows(values_only=True):
            out["table"].append(list(row))
    if "header" in wb.sheetnames:
        for row in wb["header"].iter_rows(values_only=True):
            cells = [c for c in row]
            if not any(c is not None and str(c).strip() for c in cells):
                continue
            label = cells[0] if cells else None
            value = cells[1] if len(cells) > 1 else None
            out["header"].append((label, value))
    # tolerate single-sheet grids (no named sheets): treat sheet 1 as table
    if not out["table"] and not out["header"] and wb.sheetnames:
        for row in wb[wb.sheetnames[0]].iter_rows(values_only=True):
            out["table"].append(list(row))
    return out


def _pad(row: list, n: int) -> list:
    return list(row) + [None] * (n - len(row)) if len(row) < n else list(row)[:n]


# ── positional comparison ─────────────────────────────────────────
def compare_grids(golden: dict, candidate: dict, col_labels: list[str] | None = None) -> dict:
    """Positional scores. Shapes are taken from the GOLDEN (the template)."""
    g_tab, c_tab = golden["table"], candidate["table"]
    if not g_tab:
        raise ValueError("golden has no table sheet")
    ncols = max(len(r) for r in g_tab)
    labels = col_labels or [str(v or f"col{i}") for i, v in enumerate(_pad(g_tab[0], ncols))]

    g_rows = [_pad(r, ncols) for r in g_tab[1:]]           # data rows only
    c_rows = [_pad(r, ncols) for r in c_tab[1:]] if c_tab else []
    # candidate shorter than template -> missing rows count as empty
    while len(c_rows) < len(g_rows):
        c_rows.append([None] * ncols)
    extra_rows = max(0, len(c_rows) - len(g_rows))

    n = dict(total=0, match=0, filled=0, filled_match=0, blank=0, blank_match=0,
             halluc=0, miss=0, wrong=0)
    per_col = {lab: dict(total=0, match=0, filled=0, filled_match=0)
               for lab in labels}
    cellmap = []                                            # per-cell verdicts
    for ri, (gr, cr) in enumerate(zip(g_rows, c_rows)):
        for ci in range(ncols):
            g, c = norm(gr[ci]), norm(cr[ci])
            ok = g == c
            n["total"] += 1
            n["match"] += ok
            pc = per_col[labels[ci]]
            pc["total"] += 1
            pc["match"] += ok
            verdict = "ok"
            if g:
                n["filled"] += 1
                pc["filled"] += 1
                n["filled_match"] += ok
                pc["filled_match"] += ok
                if not ok:
                    verdict = "miss" if not c else "wrong"
                    n["miss" if not c else "wrong"] += 1
            else:
                n["blank"] += 1
                n["blank_match"] += ok
                if not ok:
                    verdict = "halluc"
                    n["halluc"] += 1
            if verdict != "ok" or g:
                cellmap.append({"row": ri, "col": ci, "golden": g, "got": c,
                                "verdict": verdict})

    # header fields, matched by label (case-insensitive)
    gh = {norm(l): norm(v) for l, v in golden.get("header", []) if norm(l)}
    ch = {norm(l): norm(v) for l, v in candidate.get("header", []) if norm(l)}
    h_match = sum(1 for k, v in gh.items() if ch.get(k, "") == v)

    def frac(a, b):
        return round(a / b, 4) if b else 1.0

    res = {
        "cell_acc": frac(n["match"], n["total"]),
        "filled_recall": frac(n["filled_match"], n["filled"]),
        "blank_acc": frac(n["blank_match"], n["blank"]),
        "halluc_rate": frac(n["halluc"], n["blank"]) if n["blank"] else 0.0,
        "miss_rate": frac(n["miss"], n["filled"]) if n["filled"] else 0.0,
        "wrong_rate": frac(n["wrong"], n["filled"]) if n["filled"] else 0.0,
        "header_acc": frac(h_match, len(gh)),
        "extra_rows": extra_rows,
        "counts": n,
        "per_col": {k: {"cell_acc": frac(v["match"], v["total"]),
                        "filled_recall": frac(v["filled_match"], v["filled"]),
                        "filled": v["filled"]}
                    for k, v in per_col.items()},
        "cells": cellmap,
    }
    return res


# ── multiset comparability metric (secondary — do not steer on it) ─
def _atoms(grid: dict):
    nums, words, codes = [], [], []
    vals = [v for r in grid["table"] for v in r] + \
           [v for lv in grid.get("header", []) for v in lv]
    for v in vals:
        if v is None:
            continue
        for tok in _TOK.split(str(v).casefold()):
            tok = tok.strip(".")
            if not tok:
                continue
            if _NUM.match(tok.replace(",", ".")):
                nums.append(norm(tok))
            elif len(tok) == 1 and tok.isalpha():
                codes.append(tok)
            elif len(tok) >= 2:
                words.append(tok)
    return nums, words, codes


def _prf(g: list, c: list) -> dict:
    gc, cc = Counter(g), Counter(c)
    m = sum(min(cc[k], v) for k, v in gc.items())
    r = m / sum(gc.values()) if gc else 1.0
    p = m / sum(cc.values()) if cc else 1.0
    f = 2 * r * p / (r + p) if (r + p) else 0.0
    return {"recall": round(r, 3), "precision": round(p, 3), "f1": round(f, 3),
            "golden_total": sum(gc.values())}


def multiset_scores(golden: dict, candidate: dict) -> dict:
    gn, gw, gk = _atoms(golden)
    cn, cw, ck = _atoms(candidate)
    return {"num": _prf(gn, cn), "word": _prf(gw, cw), "code": _prf(gk, ck),
            "all": _prf(gn + gw + gk, cn + cw + ck)}


def score_files(golden_path, candidate_path) -> dict:
    g, c = load_grid(golden_path), load_grid(candidate_path)
    res = compare_grids(g, c)
    res["multiset"] = multiset_scores(g, c)
    return res


def brief_report(res: dict) -> str:
    m = res["multiset"]["all"]
    lines = [
        f"cell_acc {res['cell_acc']:.3f}  (filled_recall {res['filled_recall']:.3f}, "
        f"blank_acc {res['blank_acc']:.3f})",
        f"errors: wrong {res['wrong_rate']:.3f}  miss {res['miss_rate']:.3f}  "
        f"halluc {res['halluc_rate']:.3f}  extra_rows {res['extra_rows']}",
        f"header_acc {res['header_acc']:.3f}",
        f"multiset all-F1 {m['f1']:.3f} (secondary — position-blind)",
        "per-column cell_acc: " + "  ".join(
            f"{k}={v['cell_acc']:.2f}" for k, v in res["per_col"].items()),
    ]
    return "\n".join(lines)


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) != 2:
        print(__doc__)
        sys.exit(2)
    r = score_files(args[0], args[1])
    if "--json" in sys.argv:
        print(json.dumps(r, indent=2))
    else:
        print(brief_report(r))
