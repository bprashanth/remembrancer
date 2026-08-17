#!/usr/bin/env python3
"""Template descriptor — the single source of truth for a paperroast form.

A template is designed in the builder, so we KNOW, exactly and forever:
  * the printed column labels and their order,
  * each column's value type and (for codes/choices) its closed value domain,
  * the grid geometry in PDF points (recorded at render time, not inferred),
  * the number of data rows,
  * the form ID printed on the page (text + QR).

Everything downstream — filling, extraction, scoring, the spreadsheet the
reviewer sees — is expressed against this descriptor, which is what makes the
output isomorphic to the paper form.
"""
from __future__ import annotations

import json
import re
import secrets
import time
from dataclasses import dataclass, field, asdict
from pathlib import Path

ROOT = Path(__file__).parent.parent
TEMPLATES_DIR = ROOT / "data" / "templates"

# Column value types. "serial" is PRE-PRINTED on the form (1..N), never
# handwritten. Domains: choice/yn are closed sets; int/dec carry min/max.
COL_TYPES = ("serial", "int", "dec", "date", "yn", "choice", "name", "text")

# sensible relative widths per type when the designer does not specify
DEFAULT_WEIGHT = {"serial": 0.9, "int": 1.3, "dec": 1.5, "date": 1.9,
                  "yn": 1.0, "choice": 1.3, "name": 3.2, "text": 3.0}
HEADER_TYPES = ("text", "date", "int", "dec", "name", "choice")


@dataclass
class Column:
    key: str
    label: str
    type: str = "text"
    domain: list = field(default_factory=list)   # for choice: allowed values
    min: float | None = None                     # for int/dec
    max: float | None = None
    width: float | None = None                   # relative weight, optional
    legend: str = ""                             # e.g. "N=Normal, S=Severe"


@dataclass
class HeaderField:
    key: str
    label: str
    type: str = "text"
    domain: list = field(default_factory=list)


@dataclass
class Template:
    form_id: str
    title: str
    columns: list                                # list[Column]
    header_fields: list = field(default_factory=list)
    rows: int = 15
    orientation: str = "portrait"                # or "landscape"
    row_height: float = 26.0                     # pt; generous for handwriting
    notes: str = ""
    created: float = 0.0
    version: int = 1
    geometry: dict = field(default_factory=dict) # written by render.py

    def col(self, key):
        return next(c for c in self.columns if c.key == key)


_ID_ALPHABET = "ABCDEFGHJKMNPQRSTUVWXYZ23456789"   # no 0/O/1/I/L


def new_form_id() -> str:
    return "PR" + "".join(secrets.choice(_ID_ALPHABET) for _ in range(4))


def _slug(s: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "_", s.strip().lower()).strip("_")
    return s or "field"


def from_spec(spec: dict) -> Template:
    """Build a Template from a designer-supplied spec dict (the UI's JSON).

    Validates types and fills defaults; raises ValueError with a readable
    message on bad input."""
    title = (spec.get("title") or "").strip()
    if not title:
        raise ValueError("form needs a title")
    cols_in = spec.get("columns") or []
    if not cols_in:
        raise ValueError("form needs at least one column")
    cols, seen = [], set()
    for i, c in enumerate(cols_in):
        if isinstance(c, str):
            c = {"label": c}
        label = (c.get("label") or "").strip()
        if not label:
            raise ValueError(f"column {i+1} needs a label")
        typ = (c.get("type") or "text").strip().lower()
        if typ not in COL_TYPES:
            raise ValueError(f"column '{label}': unknown type '{typ}' "
                             f"(allowed: {', '.join(COL_TYPES)})")
        key = _slug(c.get("key") or label)
        while key in seen:
            key += "_"
        seen.add(key)
        domain = [str(d).strip() for d in (c.get("domain") or []) if str(d).strip()]
        if typ == "yn":
            domain = ["Y", "N"]
        if typ == "choice" and not domain:
            raise ValueError(f"column '{label}' is a choice column but has no "
                             f"allowed values — list them (e.g. A,B,C)")
        legend = c.get("legend") or (
            ", ".join(domain) if typ == "choice" and all(len(d) <= 3 for d in domain)
            else "")
        cols.append(Column(key=key, label=label, type=typ, domain=domain,
                           min=c.get("min"), max=c.get("max"),
                           width=c.get("width"), legend=legend))
    hdr = []
    for h in (spec.get("header_fields") or []):
        if isinstance(h, str):
            h = {"label": h}
        label = (h.get("label") or "").strip()
        if not label:
            continue
        typ = (h.get("type") or "text").strip().lower()
        if typ not in HEADER_TYPES:
            typ = "text"
        hdr.append(HeaderField(key=_slug(h.get("key") or label), label=label,
                               type=typ,
                               domain=[str(d) for d in (h.get("domain") or [])]))
    rows = int(spec.get("rows") or 15)
    if not 1 <= rows <= 60:
        raise ValueError("rows must be between 1 and 60")
    orient = spec.get("orientation") or ("landscape" if len(cols) > 7 else "portrait")
    if orient not in ("portrait", "landscape"):
        raise ValueError("orientation must be portrait or landscape")
    return Template(
        form_id=spec.get("form_id") or new_form_id(),
        title=title, columns=cols, header_fields=hdr, rows=rows,
        orientation=orient,
        row_height=float(spec.get("row_height") or 26.0),
        notes=(spec.get("notes") or "").strip(),
        created=spec.get("created") or time.time(),
        version=int(spec.get("version") or 1))


# ── persistence ───────────────────────────────────────────────────
def to_dict(t: Template) -> dict:
    d = asdict(t)
    return d


def from_dict(d: dict) -> Template:
    t = Template(**{**d, "columns": [], "header_fields": []})
    t.columns = [Column(**c) for c in d.get("columns", [])]
    t.header_fields = [HeaderField(**h) for h in d.get("header_fields", [])]
    return t


def save(t: Template, root: Path = None) -> Path:
    root = root or TEMPLATES_DIR
    d = root / t.form_id
    d.mkdir(parents=True, exist_ok=True)
    (d / "template.json").write_text(json.dumps(to_dict(t), indent=2))
    return d


def load(form_id: str, root: Path = None) -> Template:
    root = root or TEMPLATES_DIR
    p = root / form_id / "template.json"
    if not p.exists():
        raise FileNotFoundError(f"no template {form_id}")
    return from_dict(json.loads(p.read_text()))


def list_templates(root: Path = None) -> list[dict]:
    root = root or TEMPLATES_DIR
    out = []
    if root.exists():
        for p in sorted(root.glob("*/template.json")):
            try:
                d = json.loads(p.read_text())
                out.append({"form_id": d["form_id"], "title": d["title"],
                            "columns": len(d.get("columns", [])),
                            "rows": d.get("rows"), "created": d.get("created")})
            except Exception:
                continue
    return out


# ── the empty grid (answer-sheet shape) ───────────────────────────
def empty_grid(t: Template) -> dict:
    """The grid every extraction/golden is written into: header labels row +
    N empty data rows. Same shape as the paper."""
    ncols = len(t.columns)
    table = [[c.label for c in t.columns]]
    for r in range(t.rows):
        row = [None] * ncols
        for ci, c in enumerate(t.columns):
            if c.type == "serial":
                row[ci] = r + 1                     # pre-printed on the form
        table.append(row)
    return {"table": table,
            "header": [(h.label, None) for h in t.header_fields]}


def grid_to_xlsx(grid: dict, path: Path):
    import openpyxl
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("table")
    for row in grid["table"]:
        ws.append(list(row))
    wh = wb.create_sheet("header")
    for label, value in grid.get("header", []):
        wh.append([label, value])
    wb.save(str(path))
    return path
