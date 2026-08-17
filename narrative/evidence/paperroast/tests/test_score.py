#!/usr/bin/env python3
"""Scorer tests. The critical one reproduces BRIEF §2: token-accurate output
that is positionally scrambled must score LOW on cell_acc even while the
multiset F1 stays high — that divergence is the whole reason this scorer
exists."""
import sys
import tempfile
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent.parent))
from core.score import compare_grids, load_grid, multiset_scores, norm  # noqa: E402


def _grid(table, header=()):
    return {"table": table, "header": list(header)}


def test_norm():
    assert norm(" 23.0 ") == "23"
    assert norm("23,5") == norm("23.5")
    assert norm(None) == ""
    assert norm("  ") == ""
    assert norm("Coburn") == "coburn"
    assert norm("T  S") == "t s"
    assert norm(0) == "0"
    assert norm("0.0") == "0"


def test_perfect_match():
    g = _grid([["S.No", "Name", "Habit"], [1, "Coburn", "T"], [2, "Coag", "S"]],
              [("Date", "19/02/2025")])
    r = compare_grids(g, g)
    assert r["cell_acc"] == 1.0
    assert r["filled_recall"] == 1.0
    assert r["header_acc"] == 1.0
    assert r["halluc_rate"] == 0.0


def test_blanks_preserved():
    g = _grid([["A", "B"], ["x", None], [None, None]])
    c = _grid([["A", "B"], ["x", "junk"], [None, "junk2"]])
    r = compare_grids(g, c)
    # 1 filled cell correct; 3 blanks, 2 hallucinated
    assert r["filled_recall"] == 1.0
    assert abs(r["halluc_rate"] - 2 / 3) < 1e-3
    assert r["cell_acc"] == 0.5


def test_brief_linearised_failure():
    """BRIEF §2: right values, wrong cells. Multiset stays high, cell_acc low."""
    header = ["S.No", "SPP Name/Local Name", "Habit", "DBH in cms",
              "Phenological condition"]
    g = _grid(
        [header,
         [1, "Coburn", "T", 23, "Fruity"],
         [2, "Coag", "T", 9, "Flowering"]],
        [("Date", "19/02/2025"), ("Area Name", "BKM")])
    # the model's linearised emission, forced into the grid naively:
    # dropped 2 header columns, shifted cells left, stray values
    c = _grid(
        [["S.No", "SPP Name/Local Name", "Phenological condition", 1, None],
         ["Coburn", "T", 23, "Fruity", None],
         ["Coag", "T", 9, None, None]],
        [("Date", "19/02/2025"), ("Area Name", "BKM")])
    r = compare_grids(g, c)
    ms = multiset_scores(g, c)
    assert ms["all"]["f1"] > 0.75, f"multiset should look deceptively good: {ms['all']}"
    assert r["cell_acc"] < 0.15, f"positional must expose the failure: {r['cell_acc']}"
    assert r["header_acc"] == 1.0


def test_missing_rows_count_as_empty():
    g = _grid([["A"], ["x"], ["y"], ["z"]])
    c = _grid([["A"], ["x"]])
    r = compare_grids(g, c)
    assert abs(r["filled_recall"] - 1 / 3) < 1e-3
    assert r["miss_rate"] > 0.6


def test_extra_rows_flagged():
    g = _grid([["A"], ["x"]])
    c = _grid([["A"], ["x"], ["inv1"], ["inv2"]])
    r = compare_grids(g, c)
    assert r["extra_rows"] == 2


def test_xlsx_roundtrip(tmp_path=None):
    d = Path(tempfile.mkdtemp())
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    t = wb.create_sheet("table")
    t.append(["S.No", "Name"]); t.append([1, "Asha"]); t.append([2, None])
    h = wb.create_sheet("header")
    h.append(["Date", "01-02-2025"])
    p = d / "g.xlsx"; wb.save(p)
    g = load_grid(p)
    assert g["table"][1][1] == "Asha"
    assert g["header"] == [("Date", "01-02-2025")]
    r = compare_grids(g, g)
    assert r["cell_acc"] == 1.0


if __name__ == "__main__":
    fails = 0
    for name, fn in sorted({k: v for k, v in globals().items()
                            if k.startswith("test_")}.items()):
        try:
            fn()
            print(f"  ok  {name}")
        except AssertionError as e:
            fails += 1
            print(f"FAIL  {name}: {e}")
    sys.exit(1 if fails else 0)
